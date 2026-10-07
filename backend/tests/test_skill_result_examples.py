"""Ordinary skill packages run independently of the harness and report helpers."""

import json
import os
import runpy
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import jsonschema
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
PACKAGES = ROOT / "examples" / "skills"


def run_example(tmp_path, package, payload):
    source = tmp_path / "input.json"
    raw = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    source.write_bytes(raw)
    output = tmp_path / "outputs"
    result = subprocess.run(
        [
            os.sys.executable,
            "-B",
            str(PACKAGES / package / "scripts" / "build.py"),
            "--input",
            str(source),
            "--output",
            str(output),
        ],
        cwd=tmp_path,
        env={
            "PATH": os.environ.get("PATH", ""),
            "TMPDIR": str(tmp_path),
            "PYTHONIOENCODING": "utf-8",
        },
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert result.returncode == 0, result.stderr
    metadata = json.loads(result.stdout)
    files = [output / path for path in metadata["files"]]
    view_path = next(path for path in files if path.name.endswith(".view.json"))
    view = json.loads(view_path.read_text(encoding="utf-8"))
    assert view["format"] == "hartmesh.artifact-view" and view["version"] == 1
    jsonschema.Draft202012Validator(json.loads((ROOT / "contracts/artifact_view/view.schema.json").read_text(encoding="utf-8"))).validate(view)
    assert (view_path.parent / view["primary_source"]["path"]).read_bytes() == raw
    assert all((view_path.parent / item["path"]).is_file() for item in view["exports"])
    return view_path.parent, view, files


def quotes(**overrides):
    return {
        "title": "Supplier comparison",
        "suppliers": [
            {
                "name": "Supplier A",
                "quantity": "2.5",
                "unit_price": "10.01",
                "currency": "USD",
                "delivery": "14 days",
                "caveats": ["Shipping excluded"],
            },
            {
                "name": "Supplier B",
                "quantity": "2.5",
                "unit_price": "9.99",
                "currency": "USD",
                "delivery": "7 days",
                "caveats": ["Tax not supplied"],
            },
        ],
        **overrides,
    }


def procedure(**overrides):
    return {
        "title": "Opening procedure",
        "scope": "Supplied handbook only",
        "steps": [
            "Inspect the work area.",
            "Complete documented checks.",
            "Record the inspection.",
        ],
        "caveats": ["Inspection interval not supplied"],
        **overrides,
    }


def test_supplier_exact_totals_caveats_and_independent_workbook(tmp_path):
    root, view, _ = run_example(tmp_path, "supplier-comparison", quotes())
    table = next(block for block in view["blocks"] if block["type"] == "table")
    assert "25.03" in json.dumps(table) and "24.98" in json.dumps(table)
    assert "Shipping excluded" in json.dumps(view)
    workbook = load_workbook(root / "comparison.xlsx", data_only=False)
    values = [cell.value for row in workbook.active for cell in row]
    assert "25.03" in values and "24.98" in values
    assert (root / "comparison.pdf").read_bytes().startswith(b"%PDF-")


def test_procedure_preserves_text_xml_whitespace_and_caveats(tmp_path):
    text = "  Café <scope> & operations\nKeep these spaces  "
    root, view, _ = run_example(tmp_path, "procedure-summary", procedure(scope=text))
    with zipfile.ZipFile(root / "procedure.docx") as archive:
        assert archive.testzip() is None
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        parts = [node.text or "" for node in document.iter() if node.tag.endswith("}t")]
        assert "  Café <scope> & operations" in parts and "Keep these spaces  " in parts
        assert "Inspection interval not supplied" in parts
        assert all('TargetMode="External"' not in archive.read(name).decode("utf-8") for name in archive.namelist() if name.endswith(".rels"))
    assert next(block for block in view["blocks"] if block["type"] == "list")["items"] == procedure()["steps"]


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_workbook_labels_remain_exact_inert_strings(tmp_path, prefix):
    payload = quotes()
    name = prefix + "literal supplier"
    payload["suppliers"][0]["name"] = name
    root, _, _ = run_example(tmp_path, "supplier-comparison", payload)
    workbook = load_workbook(root / "comparison.xlsx", data_only=False)
    cell = next(cell for row in workbook.active for cell in row if cell.value == name)
    assert cell.data_type == "s"


@pytest.mark.parametrize(
    "package,payload,export",
    [
        ("supplier-comparison", quotes(title="Résumé 📋"), "comparison.xlsx"),
        ("procedure-summary", procedure(title="Procedure 📋"), "procedure.docx"),
    ],
)
def test_unsupported_portable_pdf_text_preserves_source_editable_export_and_explicit_notice(tmp_path, package, payload, export):
    root, view, _ = run_example(tmp_path, package, payload)
    assert (root / export).is_file()
    assert not any(item["path"].endswith(".pdf") for item in view["exports"])
    assert any(block["type"] == "notice" and "PDF" in block["text"] for block in view["blocks"])


@pytest.mark.parametrize("amount", [True, "NaN", "Infinity", "1e999", "1.123456", "9" * 40])
def test_invalid_price_is_rejected_before_publishing_any_outputs(tmp_path, amount):
    payload = quotes()
    payload["suppliers"][0]["unit_price"] = amount
    source = tmp_path / "input.json"
    source.write_text(json.dumps(payload), encoding="utf-8")
    output = tmp_path / "outputs"
    result = subprocess.run(
        [
            os.sys.executable,
            "-B",
            str(PACKAGES / "supplier-comparison" / "scripts" / "build.py"),
            "--input",
            str(source),
            "--output",
            str(output),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert result.returncode != 0
    assert "Invalid unit_price" in result.stderr
    assert not output.exists() or not list(output.iterdir())


def test_maximum_quotations_preserve_all_supplied_caveats_in_valid_bounded_lists(
    tmp_path,
):
    payload = quotes()
    payload["suppliers"] = [
        {
            "name": f"Supplier {index}",
            "quantity": "1",
            "unit_price": "1.00",
            "currency": "USD",
            "caveats": [f"Caveat {item}" for item in range(16)],
        }
        for index in range(50)
    ]
    _, view, _ = run_example(tmp_path, "supplier-comparison", payload)
    caveats = [item for block in view["blocks"] if block["type"] == "list" for item in block["items"]]
    assert len(caveats) == 800


def test_maximum_procedure_long_words_omit_pdf_explicitly_without_clipping_editable_content(
    tmp_path,
):
    steps = ["x" * 2048 for _ in range(24)]
    root, view, _ = run_example(tmp_path, "procedure-summary", procedure(steps=steps))
    assert next(block for block in view["blocks"] if block["type"] == "list")["items"] == steps
    assert not (root / "procedure.pdf").exists()
    assert any(block["type"] == "notice" and "eight-page" in block["text"] for block in view["blocks"])
    with zipfile.ZipFile(root / "procedure.docx") as archive:
        assert archive.read("word/document.xml").count(steps[0].encode()) == 24


@pytest.mark.parametrize("package", ["supplier-comparison", "procedure-summary"])
@pytest.mark.parametrize(
    "raw,error",
    [
        (b'{"title":"a","title":"b"}', "Duplicate input key"),
        (b'{"value":NaN}', "Non-finite input number"),
        (b" " * (256 * 1024 + 1), "Input exceeds 256 KiB"),
        (b"[]", "Input must be an object"),
    ],
)
def test_input_bounds_before_document_writers(tmp_path, package, raw, error):
    module = runpy.run_path(str(PACKAGES / package / "scripts/documents.py"))
    source = tmp_path / "input.json"
    source.write_bytes(raw)
    with pytest.raises(ValueError, match=error):
        module["read_input"](source)


@pytest.mark.parametrize("package", ["supplier-comparison", "procedure-summary"])
def test_format_and_view_failures_preserve_complete_ordinary_files(tmp_path, package):
    module = runpy.run_path(str(PACKAGES / package / "scripts/documents.py"))

    def failed(path):
        path.write_bytes(b"partial")
        raise RuntimeError("fixture writer failure")

    result = module["publish"](
        tmp_path / "outputs",
        "example",
        b'{"original":"exact"}',
        {"title": "Fixture", "blocks": [], "unserializable": object()},
        [
            ("failed.pdf", "PDF", failed),
            ("successful.txt", "Text", lambda path: path.write_bytes(b"complete")),
        ],
    )
    files = [tmp_path / "outputs" / path for path in result["files"]]
    assert {path.name for path in files} == {"example.json", "successful.txt"}
    assert next(path for path in files if path.suffix == ".json").read_bytes() == b'{"original":"exact"}'
    assert next(path for path in files if path.suffix == ".txt").read_bytes() == b"complete"
    assert len(result["warnings"]) == 2
    assert not list((tmp_path / "outputs").rglob("failed.pdf"))
    assert not list((tmp_path / "outputs").glob(".skill-result-*"))


def test_maximum_decimal_product_is_exact(tmp_path):
    payload = quotes()
    payload["suppliers"] = [dict(payload["suppliers"][0], quantity="1000000.000", unit_price="1000000000.00")]
    _, view, _ = run_example(tmp_path, "supplier-comparison", payload)
    table = next(block for block in view["blocks"] if block["type"] == "table")
    assert table["rows"][0][3] == "1000000000000000.00"


@pytest.mark.parametrize("character", ["\ufffe", "\uffff"])
def test_xml_incompatible_procedure_text_preserves_source_and_view_without_invalid_docx(tmp_path, character):
    _, view, files = run_example(tmp_path, "procedure-summary", procedure(scope="Exact " + character))
    assert all(path.suffix != ".docx" for path in files)
    assert any(block["type"] == "notice" and "XML-incompatible" in block["text"] for block in view["blocks"])
    assert next(block for block in view["blocks"] if block["type"] == "text")["paragraphs"] == ["Exact " + character]


@pytest.mark.parametrize("package,editable", [("supplier-comparison", "comparison.xlsx"), ("procedure-summary", "procedure.docx")])
@pytest.mark.parametrize("character,reason", [("\x7f", "DEL"), ("\u00ad", "soft hyphen")])
def test_undefined_pdf_control_glyph_omits_only_pdf_preserving_exact_title(tmp_path, package, editable, character, reason):
    title = "Before" + character + "After"
    payload = quotes(title=title) if package == "supplier-comparison" else procedure(title=title)
    root, view, files = run_example(tmp_path, package, payload)
    assert view["title"] == title
    assert (root / editable).is_file()
    assert all(path.suffix != ".pdf" for path in files)
    assert any(block["type"] == "notice" and reason in block["text"] for block in view["blocks"])
    if editable.endswith(".xlsx"):
        workbook = load_workbook(root / editable)
        assert title in [cell.value for row in workbook.active for cell in row]

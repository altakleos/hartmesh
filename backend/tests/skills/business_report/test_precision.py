"""Source-precision amounts survive JSON and live/cached XLSX calculations."""

from __future__ import annotations

import csv
import importlib.util
import json
import re
import sys
import xml.etree.ElementTree as ET
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

for _library in ("pandas", "docx", "jinja2", "xlsxwriter"):
    pytest.importorskip(_library)

SCRIPT = Path(__file__).resolve().parents[4] / "skills/public/business-report/scripts/report.py"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@pytest.fixture(scope="module")
def report_script():
    spec = importlib.util.spec_from_file_location("business_report_precision_script", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


def _build_workbook(report_script, tmp_path, amounts, *, distinct_groups):
    source = tmp_path / "amounts.csv"
    with source.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Date", "Amount", "Technician", "Service", "Customer", "Status", "Source"])
        for index, amount in enumerate(amounts):
            group = f"Group {index}" if distinct_groups else "One group"
            writer.writerow(
                [
                    f"2026-09-{index % 30 + 1:02d}",
                    amount,
                    group,
                    group,
                    group,
                    "Unpaid",
                    group,
                ]
            )
    original = source.read_bytes()
    with source.open(newline="", encoding="utf-8") as stream:
        expected = [Decimal(row["Amount"]) for row in csv.DictReader(stream)]
    ctx = report_script.prepare(
        [str(source)],
        "2026-09",
        report_script.BuildOptions(),
        None,
        report_script.load_profile(None, None),
    )
    built, _ = report_script.build_report(ctx, 0, report_script.compute_checks)
    # The renderer consumes the JSON document, including its numeric precision.
    report_path = tmp_path / "precision.report.json"
    report_path.write_text(json.dumps(built), encoding="utf-8")
    built = json.loads(report_path.read_text(encoding="utf-8"))
    xlsx = report_script.render(built, report_path, "xlsx", None, None)
    assert source.read_bytes() == original
    return built, xlsx, expected


def _evaluate_formula(workbook, sheet_name, coordinate):
    """Evaluate the emitted SUM/COUNTA/ratio formulas from workbook operands."""
    cell = workbook[sheet_name][coordinate]
    value = cell.value
    if cell.data_type != "f":
        return Decimal(str(value or 0))
    if match := re.fullmatch(r"=(SUM|COUNTA)\((?:(\w+)!)?([A-Z]+)(\d+):([A-Z]+)(\d+)\)", value):
        operation, referenced_sheet, first_column, first_row, last_column, last_row = match.groups()
        assert first_column == last_column
        referenced_sheet = referenced_sheet or sheet_name
        cells = [f"{first_column}{row}" for row in range(int(first_row), int(last_row) + 1)]
        if operation == "COUNTA":
            return Decimal(sum(workbook[referenced_sheet][address].value is not None for address in cells))
        return sum(
            (_evaluate_formula(workbook, referenced_sheet, address) for address in cells),
            Decimal(0),
        )
    match = re.fullmatch(r"=IF\(([A-Z]+\d+)=0,0,([A-Z]+\d+)/([A-Z]+\d+)\)", value)
    assert match, value
    denominator, numerator, repeated_denominator = match.groups()
    assert denominator == repeated_denominator
    divisor = _evaluate_formula(workbook, sheet_name, denominator)
    return _evaluate_formula(workbook, sheet_name, numerator) / divisor if divisor else Decimal(0)


def _assert_formula_caches_match_operands(xlsx):
    formulas = openpyxl.load_workbook(xlsx, data_only=False)
    cached = openpyxl.load_workbook(xlsx, data_only=True)
    count = 0
    with ZipFile(xlsx) as archive:
        for index, sheet in enumerate(formulas, start=1):
            xml = ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            for row in sheet:
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    raw = xml.find(f'.//s:c[@r="{cell.coordinate}"]', NS)
                    assert raw.find("s:f", NS).text == cell.value[1:]
                    raw_cached = Decimal(raw.find("s:v", NS).text)
                    expected = _evaluate_formula(formulas, sheet.title, cell.coordinate)
                    # XLSX numbers are binary floats; reject cent-level drift while
                    # allowing representation error on repeating average ratios.
                    assert float(raw_cached) == pytest.approx(float(expected), rel=1e-13, abs=1e-13), (sheet.title, cell.coordinate)
                    assert cached[sheet.title][cell.coordinate].value == float(raw_cached)
                    count += 1
    assert count >= 3
    return formulas, cached


@pytest.mark.parametrize("amounts", [["1.005"] * 5, ["2.675", "1.005", "-0.335", "0.0005", "0.00"]])
def test_rows_and_summary_preserve_amount_precision(report_script, tmp_path, amounts):
    built, xlsx, expected = _build_workbook(report_script, tmp_path, amounts, distinct_groups=True)
    workbook, cached = _assert_formula_caches_match_operands(xlsx)
    assert workbook["Summary"]["B2"].value == f"=SUM(Rows!B2:B{len(expected) + 1})"
    assert [Decimal(str(workbook["Rows"].cell(index + 2, 2).value)) for index in range(len(expected))] == expected
    assert [Decimal(str(row[1])) for row in built["rows"]["rows"]] == expected
    total = sum(expected, Decimal(0))
    assert Decimal(str(built["kpis"][0]["value"])) == total
    assert Decimal(str(cached["Summary"]["B2"].value)) == total
    assert cached["Summary"]["B4"].value == pytest.approx(float(total / len(expected)))
    assert workbook["Rows"]["B2"].number_format == '"$"#,##0.00'
    assert workbook["Summary"]["B2"].number_format == '"$"#,##0.00'
    check = next(check for check in built["checks"] if check["id"] == "totals_reconcile")
    assert check["status"] == "pass"
    displayed = f"${total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):,.2f}"
    assert displayed in check["text"]
    html = report_script.render(built, tmp_path / "precision.report.json", "html", None, None)
    assert f'<span class="kpi-value">{displayed}</span>' in html.read_text(encoding="utf-8")


@pytest.mark.parametrize("distinct_groups,count", [(False, 5), (True, 30)])
def test_group_totals_unpaid_and_other_rows_keep_amount_precision(report_script, tmp_path, distinct_groups, count):
    built, xlsx, expected = _build_workbook(report_script, tmp_path, ["1.005"] * count, distinct_groups=distinct_groups)
    _assert_formula_caches_match_operands(xlsx)
    total = sum(expected, Decimal(0))
    assert Decimal(str(next(kpi for kpi in built["kpis"] if kpi["id"] == "unpaid")["value"])) == total
    for section in built["sections"]:
        table = section.get("table")
        if not table or not table.get("totals"):
            continue
        for heading in ("Revenue", "Unpaid"):
            if heading not in table["columns"]:
                continue
            column = table["columns"].index(heading)
            assert sum((Decimal(str(row[column])) for row in table["rows"]), Decimal(0)) == total, section["id"]
            assert Decimal(str(table["totals"][column])) == total, section["id"]
        if distinct_groups and section["id"] == "by_person":
            assert table["rows"][-1][0] == "Other (5 technicians)"
            assert Decimal(str(table["rows"][-1][2])) == Decimal("5.025")

"""Run inside an existing AIO image with read-only skill mounts and writable output."""

from __future__ import annotations

import errno
import hashlib
import json
import subprocess
from pathlib import Path

import docx
import fitz
from openpyxl import load_workbook


def main():
    output = Path("/mnt/user-data/outputs")
    receipts = []
    for package, mounted, filename, stem in (
        (
            "supplier-comparison",
            "/mnt/skills/public/supplier-comparison",
            "quotes.json",
            "comparison",
        ),
        (
            "procedure-summary",
            "/mnt/skills/custom/private-procedure",
            "procedure.json",
            "procedure",
        ),
    ):
        root = Path(mounted)
        for path in (root / "SKILL.md", root / "scripts/build.py"):
            original = path.read_bytes()
            try:
                path.write_bytes(original + b"fixture mutation")
            except OSError as error:
                assert error.errno == errno.EROFS, error
            else:
                raise AssertionError("skill mount permitted a write")
            assert path.read_bytes() == original
        source = Path("/mnt/inputs") / filename
        result = subprocess.run(
            [
                "python3",
                "-B",
                str(root / "scripts/build.py"),
                "--input",
                str(source),
                "--output",
                str(output),
            ],
            capture_output=True,
            text=True,
            check=True,
            timeout=30,
        )
        manifest = json.loads(result.stdout)
        paths = [output / name for name in manifest["files"]]
        assert len(paths) == 4
        bundle = paths[0].parent
        assert (bundle / (stem + ".json")).read_bytes() == source.read_bytes()
        with fitz.open(bundle / (stem + ".pdf")) as pdf:
            assert not pdf.is_repaired
            assert 1 <= len(pdf) <= 8
            text = "\n".join(page.get_text() for page in pdf)
            if stem == "comparison":
                assert "25.03" in text and "Shipping excluded" in text
            else:
                assert "Inspect the work area." in text
        if stem == "comparison":
            book = load_workbook(bundle / "comparison.xlsx")
            values = [cell.value for row in book.active for cell in row]
            assert "25.03" in values and "Supplier A: Shipping excluded" in values
        else:
            document = docx.Document(bundle / "procedure.docx")
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            assert "1. Inspect the work area." in paragraphs
            assert "Inspection interval not supplied" in paragraphs
        receipts.append(
            {
                "package": package,
                "mounted_path": mounted,
                "files": {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths},
                "readonly_writes_refused": True,
                "independent_readers": True,
            }
        )
    (output / "sandbox-result.json").write_text(json.dumps(receipts, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipts))


if __name__ == "__main__":
    main()

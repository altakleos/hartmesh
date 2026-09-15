"""Renderers for the business-report skill: HTML, PDF, DOCX, XLSX and the charts.

Every renderer reads only report.json (plus the chart PNGs and the brand logo,
both confined to the report directory and the tenant bundle). Nothing here
touches the network: the PDF printer refuses every URL that is not inline data.
"""

from __future__ import annotations

import base64
import datetime as dt
import math
import re
import sys
from pathlib import Path

import docx
import jinja2
import xlsxwriter
import xlsxwriter.utility
from business_report_common import (
    CHART_PNG_PATTERN,
    CURRENCY_SYMBOLS,
    DEFAULT_BRAND,
    EXIT_MISSING_LIBRARY,
    MISSING_LIBRARY_MESSAGE,
    REPORT_SUFFIX,
    TEMPLATES_DIR,
    InputError,
    axis_label,
    cell_format,
    checks_line,
    format_value,
    is_color,
    load_brand,
    picture_inside,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

CHART_SIZE_INCHES = (6.4, 2.9)
CHART_DPI = 200
SUM_TOLERANCE_PER_ROW = 0.005  # each rounded cell may be half a cent off its unrounded value


def data_uri(path: Path) -> str:
    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}[path.suffix.lower()]
    return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"


def chart_path(report_path: Path, png: str) -> Path | None:
    """The PNG a chart entry points at, only when it is the charts/<id>.png the schema allows."""

    if not isinstance(png, str) or not CHART_PNG_PATTERN.match(png):
        raise InputError(f"Chart path {png!r} is not of the form charts/<id>.png; the report was not produced by this skill.")
    return picture_inside(str(report_path.parent / png), [report_path.parent])


def effective_brand(report: dict, report_path: Path, tenant_dir: str | None) -> dict:
    """The brand the report was built with, the tenant bundle on top when given; pictures confined."""

    brand = dict(DEFAULT_BRAND)
    stored = report["meta"].get("brand", {})
    brand["company"] = str(stored.get("company") or report["meta"].get("company") or "")
    for key in ("primary", "secondary"):
        if is_color(stored.get(key)):
            brand[key] = stored[key]
    roots = [report_path.parent, Path(tenant_dir)] if tenant_dir else [report_path.parent]
    logo = picture_inside(stored.get("logo"), roots)
    brand["logo"] = str(logo) if logo else None
    if tenant_dir:
        tenant = load_brand(tenant_dir)
        for key in ("primary", "secondary"):
            if tenant[key] != DEFAULT_BRAND[key]:
                brand[key] = tenant[key]
        if tenant["company"] and not brand["company"]:
            brand["company"] = tenant["company"]
        if tenant["logo"]:
            brand["logo"] = tenant["logo"]
    return brand


_INLINE_FETCHER = []


def fetch_inline_only(url: str, *args, **kwargs):
    """WeasyPrint URL fetcher that serves inline data and refuses everything else, on WeasyPrint 68 and later."""

    if not url.startswith("data:"):
        raise ValueError(f"External resources are not loaded while printing a report (refused: {url[:80]}).")
    if not _INLINE_FETCHER:
        import weasyprint.urls

        _INLINE_FETCHER.append(weasyprint.urls.URLFetcher(allowed_protocols={"data"}))
    return _INLINE_FETCHER[0].fetch(url, *args, **kwargs)


fetch_inline_only._fail_on_errors = True  # a refused URL stops the render instead of printing a gap


def render_html(report: dict, report_path: Path, brand: dict) -> str:
    environment = jinja2.Environment(loader=jinja2.FileSystemLoader(str(TEMPLATES_DIR)), autoescape=jinja2.select_autoescape(["html", "j2"]), undefined=jinja2.StrictUndefined)
    template = environment.get_template("report.html.j2")
    currency = report["meta"]["currency"]["code"]
    charts = {}
    for chart in report["charts"]:
        png = chart_path(report_path, chart["png"])
        if png is not None:
            charts[chart["id"]] = data_uri(png)
    logo = data_uri(Path(brand["logo"])) if brand.get("logo") else None
    return template.render(
        report=report,
        brand=brand,
        company=brand.get("company") or report["meta"].get("company", ""),
        logo=logo,
        charts=charts,
        css=(TEMPLATES_DIR / "report.css").read_text(encoding="utf-8"),
        fmt=lambda value, fmt: format_value(value, fmt, currency),
        cell_fmt=cell_format,
        checks_line=checks_line(report),
    )


def render_pdf(html: str, path: Path) -> None:
    try:
        import weasyprint
    except ImportError as error:
        sys.stderr.write(f"{MISSING_LIBRARY_MESSAGE} ({error})\n")
        sys.exit(EXIT_MISSING_LIBRARY)
    try:
        weasyprint.HTML(string=html, url_fetcher=fetch_inline_only).write_pdf(str(path))
    except weasyprint.urls.FatalURLFetchingError as error:
        raise InputError(f"The report refers to a resource outside the report; nothing external is loaded. {error}") from error


def _fill_table(word_table, table: dict, currency: str) -> None:
    """Fill a python-docx table row by row (cell() per cell is quadratic in the table size)."""

    rows = list(table["rows"]) + ([table["totals"]] if table.get("totals") else [])
    header_cells = word_table.rows[0].cells
    for index, column in enumerate(table["columns"]):
        header_cells[index].text = column
        header_cells[index].paragraphs[0].runs[0].bold = True
    for row_index, row in enumerate(rows, start=1):
        cells = word_table.rows[row_index].cells
        is_totals = bool(table.get("totals")) and row_index == len(rows)
        for index, value in enumerate(row):
            fmt = cell_format(table, None if is_totals else row_index - 1, index)
            cells[index].text = format_value(value, fmt, currency)
            paragraph = cells[index].paragraphs[0]
            if fmt != "text":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if is_totals and paragraph.runs:
                paragraph.runs[0].bold = True


def render_docx(report: dict, report_path: Path, brand: dict, path: Path) -> None:
    currency = report["meta"]["currency"]["code"]
    document = docx.Document()
    if brand.get("logo"):
        document.add_picture(brand["logo"], width=Inches(1.5))
    document.add_heading(report["meta"]["title"], level=0)
    company = brand.get("company") or report["meta"].get("company", "")
    document.add_paragraph(" · ".join(part for part in (company, report["meta"]["period"]["label"], f"draft {report['meta']['draft']}") if part))
    kpi_table = document.add_table(rows=2, cols=max(1, len(report["kpis"])))
    kpi_table.style = "Table Grid"
    labels, values = kpi_table.rows[0].cells, kpi_table.rows[1].cells
    for index, kpi in enumerate(report["kpis"]):
        labels[index].text = kpi["label"]
        value = format_value(kpi["value"], kpi["format"], currency)
        delta = kpi.get("delta")
        if delta and delta.get("pct") is not None:
            value += f" ({'+' if delta['pct'] > 0 else ''}{format_value(delta['pct'], 'percent')} vs {delta['vs']})"
        values[index].text = value
    for section in report["sections"]:
        document.add_heading(section["heading"], level=1)
        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets", []):
            document.add_paragraph(bullet, style="List Bullet")
        table = section.get("table")
        if table:
            row_count = len(table["rows"]) + (1 if table.get("totals") else 0)
            word_table = document.add_table(rows=row_count + 1, cols=len(table["columns"]))
            word_table.style = "Table Grid"
            _fill_table(word_table, table, currency)
        for chart_id in section.get("charts", []):
            entry = next((chart for chart in report["charts"] if chart["id"] == chart_id), None)
            png = chart_path(report_path, entry["png"]) if entry else None
            if png is not None:
                document.add_picture(str(png), width=Inches(6.0))
        if section.get("note"):
            document.add_paragraph(section["note"])
    document.add_heading("Checks", level=1)
    document.add_paragraph(checks_line(report))
    for check in report["checks"]:
        document.add_paragraph(f"{check['status'].replace('_', ' ').capitalize()}: {check['text']}", style="List Bullet")
    if report["notes"]:
        document.add_paragraph("Not included")
        for note in report["notes"]:
            document.add_paragraph(note, style="List Bullet")
    inputs = ", ".join(f"{entry['name']} (uploaded {entry['uploaded']})" for entry in report["meta"]["inputs"])
    document.add_paragraph(f"Built from {inputs}. Checked by the report script.")
    document.save(str(path))


def sheet_name(heading: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\']", "", heading).strip()[:31] or "Sheet"
    candidate, counter = name, 2
    while candidate in used:
        suffix = f" {counter}"
        candidate = name[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate)
    return candidate


def render_xlsx(report: dict, path: Path) -> None:
    currency = report["meta"]["currency"]["code"]
    workbook = xlsxwriter.Workbook(str(path), {"strings_to_formulas": False, "strings_to_numbers": False, "strings_to_urls": False})
    symbol = CURRENCY_SYMBOLS.get(currency)
    money_format = f'"{symbol}"#,##0.00' if symbol else f'"{currency} "#,##0.00'
    number_formats = {"currency": money_format, "integer": "#,##0", "number": "#,##0.00", "percent": '0.0"%"', "date": "yyyy-mm-dd"}
    formats = {fmt: workbook.add_format({"num_format": pattern}) for fmt, pattern in number_formats.items()}
    formats["header"] = workbook.add_format({"bold": True, "bg_color": "#EEEEEE", "border": 1})
    formats["total"] = workbook.add_format({"bold": True, "top": 1})
    for fmt, pattern in number_formats.items():
        formats[f"total_{fmt}"] = workbook.add_format({"bold": True, "top": 1, "num_format": pattern})
    used_names: set[str] = {"Summary", "Rows"}

    def write_cell(sheet, row: int, column: int, value, fmt: str, cell_style=None) -> None:
        if fmt == "text":
            sheet.write_string(row, column, "" if value is None else str(value), cell_style)
        elif value is None or (isinstance(value, float) and math.isnan(value)):
            sheet.write_blank(row, column, None, cell_style)
        elif fmt == "date":
            try:
                parsed = dt.datetime.fromisoformat(str(value))
            except ValueError:
                sheet.write_string(row, column, str(value), cell_style)
            else:
                sheet.write_datetime(row, column, parsed, cell_style or formats["date"])
        else:
            sheet.write_number(row, column, float(value), cell_style or formats[fmt])

    def write_table(sheet, table: dict, live_totals: bool) -> None:
        for column, heading in enumerate(table["columns"]):
            sheet.write_string(0, column, heading, formats["header"])
            sheet.set_column(column, column, max(12, min(40, len(heading) + 2)))
        for row_index, row in enumerate(table["rows"], start=1):
            for column, value in enumerate(row):
                write_cell(sheet, row_index, column, value, cell_format(table, row_index - 1, column))
        if not table.get("totals"):
            sheet.freeze_panes(1, 0)
            return
        total_row = len(table["rows"]) + 1
        first_data, last_data = 2, len(table["rows"]) + 1
        ratios = {ratio[0]: (ratio[1], ratio[2]) for ratio in table.get("ratios", [])}
        for column, value in enumerate(table["totals"]):
            fmt = table["formats"][column]
            letter = xlsxwriter.utility.xl_col_to_name(column)
            column_values = [row[column] for row in table["rows"] if isinstance(row[column], (int, float)) and not isinstance(row[column], bool)]
            numeric_total = isinstance(value, (int, float)) and not isinstance(value, bool)
            reconciles = numeric_total and abs(sum(column_values) - float(value)) <= 0.01 + SUM_TOLERANCE_PER_ROW * len(table["rows"])
            if fmt == "text":
                sheet.write_string(total_row, column, "" if value is None else str(value), formats["total"])
            elif column in ratios and live_totals and table["rows"]:
                numerator, denominator = (xlsxwriter.utility.xl_col_to_name(index) + str(total_row + 1) for index in ratios[column])
                sheet.write_formula(total_row, column, f"=IF({denominator}=0,0,{numerator}/{denominator})", formats[f"total_{fmt}"], float(value or 0))
            elif live_totals and fmt in ("currency", "integer", "number") and column > 0 and table["rows"] and reconciles:
                sheet.write_formula(total_row, column, f"=SUM({letter}{first_data}:{letter}{last_data})", formats[f"total_{fmt}"], float(value or 0))
            else:
                write_cell(sheet, total_row, column, value, fmt, formats.get(f"total_{fmt}", formats["total"]))
        sheet.freeze_panes(1, 0)

    summary = workbook.add_worksheet("Summary")
    rows_table = report["rows"]
    row_count = len(rows_table["rows"])
    amount_index = rows_table["formats"].index("currency") if "currency" in rows_table["formats"] else None
    amount_letter = xlsxwriter.utility.xl_col_to_name(amount_index) if amount_index is not None else None
    summary.set_column(0, 0, 28)
    summary.set_column(1, 1, 22)
    summary.write_string(0, 0, "Metric", formats["header"])
    summary.write_string(0, 1, "Value", formats["header"])
    line = 1
    kpi_cells: dict[str, int] = {}
    for kpi in report["kpis"]:
        summary.write_string(line, 0, kpi["label"])
        if kpi["id"] == "revenue" and amount_letter and row_count:
            summary.write_formula(line, 1, f"=SUM(Rows!{amount_letter}2:{amount_letter}{row_count + 1})", formats["currency"], float(kpi["value"] or 0))
        elif kpi["id"] == "jobs" and row_count:
            summary.write_formula(line, 1, f"=COUNTA(Rows!A2:A{row_count + 1})", formats["integer"], float(kpi["value"] or 0))
        elif kpi["id"] == "average_ticket" and "revenue" in kpi_cells and "jobs" in kpi_cells:
            revenue_cell, jobs_cell = f"B{kpi_cells['revenue'] + 1}", f"B{kpi_cells['jobs'] + 1}"
            summary.write_formula(line, 1, f"=IF({jobs_cell}=0,0,{revenue_cell}/{jobs_cell})", formats["currency"], float(kpi["value"] or 0))
        else:
            write_cell(summary, line, 1, kpi["value"], kpi["format"])
        kpi_cells[kpi["id"]] = line
        line += 1
    line += 1
    meta = report["meta"]
    details = (
        ("Title", meta["title"]),
        ("Company", meta["company"]),
        ("Period", meta["period"]["label"]),
        ("Generated", meta["generated_at"]),
        ("Draft", str(meta["draft"])),
        ("Inputs", ", ".join(entry["name"] for entry in meta["inputs"])),
        ("Checks", checks_line(report)),
    )
    for label, value in details:
        summary.write_string(line, 0, label)
        summary.write_string(line, 1, value)
        line += 1
    for section in report["sections"]:
        table = section.get("table")
        if not table:
            continue
        write_table(workbook.add_worksheet(sheet_name(section["heading"], used_names)), table, live_totals=True)
    write_table(workbook.add_worksheet("Rows"), rows_table, live_totals=False)
    workbook.close()


def draw_charts(charts: dict[str, dict], out_dir: Path, brand: dict, currency: str) -> list[Path]:
    if not charts:
        return []
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.ticker import FuncFormatter
    except ImportError as error:
        sys.stderr.write(f"{MISSING_LIBRARY_MESSAGE} ({error})\n")
        sys.exit(EXIT_MISSING_LIBRARY)
    charts_dir = out_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for chart_id, data in charts.items():
        figure, axis = plt.subplots(figsize=CHART_SIZE_INCHES, dpi=CHART_DPI)
        labels, values = data["labels"], data["values"]
        formatter = FuncFormatter(lambda value, _pos, fmt=data["format"]: axis_label(value, fmt, currency))
        if data["type"] == "barh":
            axis.barh(labels[::-1], values[::-1], color=brand["primary"])
            axis.xaxis.set_major_formatter(formatter)
            axis.grid(axis="x", color="#dddddd", linewidth=0.6)
        else:
            axis.bar(labels, values, color=brand["primary"])
            axis.yaxis.set_major_formatter(formatter)
            axis.grid(axis="y", color="#dddddd", linewidth=0.6)
            if len(labels) > 6 or max((len(label) for label in labels), default=0) > 8:
                plt.setp(axis.get_xticklabels(), rotation=30, ha="right")
        axis.set_axisbelow(True)
        axis.set_title(data["title"], loc="left", fontsize=11, color="#222222")
        for side in ("top", "right"):
            axis.spines[side].set_visible(False)
        axis.tick_params(labelsize=8, colors="#333333")
        figure.tight_layout()
        path = charts_dir / f"{chart_id}.png"
        figure.savefig(path, dpi=CHART_DPI)
        plt.close(figure)
        written.append(path)
    return written


def render(report: dict, report_path: Path, target: str, out_path: Path | None, tenant_dir: str | None) -> Path:
    brand = effective_brand(report, report_path, tenant_dir)
    base = report_path.name[: -len(REPORT_SUFFIX)] if report_path.name.endswith(REPORT_SUFFIX) else report_path.stem
    path = out_path or report_path.with_name(f"{base}.{target}")
    path.parent.mkdir(parents=True, exist_ok=True)
    if target == "html":
        path.write_text(render_html(report, report_path, brand), encoding="utf-8")
    elif target == "pdf":
        render_pdf(render_html(report, report_path, brand), path)
    elif target == "docx":
        render_docx(report, report_path, brand, path)
    elif target == "xlsx":
        render_xlsx(report, path)
    else:
        raise InputError(f"Unknown render target {target!r}; use html, pdf, docx or xlsx.")
    return path
